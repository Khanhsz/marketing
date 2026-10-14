"""Slide-first case presentation: Gemini slide architecture, 16:9 preview, PPTX export."""
import html
import io
import json
import re

import streamlit as st
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

LAYOUTS = {"cover", "split", "data", "comparison", "timeline", "statement", "framework"}
ORDER = (
    "Cover / case title",
    "Business & marketing challenge",
    "Executive summary",
    "Category and market evidence",
    "Company diagnosis",
    "Competitor landscape and white space",
    "Consumer segments / purchase barriers",
    "Consumer tension and sharp insight",
    "Core strategic problem",
    "Positioning and strategic direction",
    "Big Idea, key message, creative rationale",
    "IMC journey, phases",
    "Hero activation and channel plan",
    "Budget / KPIs / measurement",
    "Feasibility, risks, recommendation and closing"
)

def parse_json(raw):
    raw = (raw or "").strip()
    mark = chr(96) * 3
    if raw.startswith(mark):
        raw = re.sub(r"^" + re.escape(mark) + r"(?:json)?\s*", "", raw, flags=re.I)
        raw = re.sub(r"\s*" + re.escape(mark) + r"$", "", raw)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        a, b = raw.find("{"), raw.rfind("}")
        if a < 0 or b < a:
            raise ValueError("Slide Architect không trả JSON hợp lệ.")
        return json.loads(raw[a:b+1])

def normalize(raw, permitted_ids, count=15):
    obj = parse_json(raw)
    if not isinstance(obj, dict) or not isinstance(obj.get("slides"), list):
        raise ValueError("Thiếu danh sách slides.")
    if len(obj["slides"]) != count:
        raise ValueError(f"Cần đúng {count} slides, nhận được {len(obj['slides'])}.")
    slides = []
    for index, raw_slide in enumerate(obj["slides"], 1):
        if not isinstance(raw_slide, dict):
            raise ValueError("Slide không phải object.")
        points = raw_slide.get("insights") or []
        if not isinstance(points, list):
            raise ValueError("insights phải là một mảng.")
        points = [str(p).strip()[:230] for p in points if str(p).strip()][:4]
        mode = str(raw_slide.get("layout") or "split").lower().strip()
        mode = mode if mode in LAYOUTS else "split"
        takeaway = str(raw_slide.get("takeaway") or "").strip()[:310]
        if mode != "cover" and (len(points) < 2 or not takeaway):
            raise ValueError(f"Slide {index} cần >=2 insights và một Key Takeaway.")
        src = raw_slide.get("source_ids") or []
        if not isinstance(src, list):
            src = []
        ids = list(dict.fromkeys(
            s for s in (str(x).strip(" []").upper() for x in src) if s in permitted_ids
        ))
        slides.append({
            "slide_no": index, "section": str(raw_slide.get("section") or "Marketing Case")[:70],
            "headline": str(raw_slide.get("headline") or "Untitled slide")[:150],
            "kicker": str(raw_slide.get("kicker") or "MARKETING CASE")[:65],
            "layout": mode, "insights": points, "takeaway": takeaway,
            "visual": str(raw_slide.get("visual") or "Gợi ý layout, chưa có biểu đồ định lượng.")[:350],
            "source_ids": ids,
            "presenter_note": str(raw_slide.get("presenter_note") or "")[:800],
        })
    return {"deck_title": str(obj.get("deck_title") or "Marketing Case Pitch")[:150],
            "deck_subtitle": str(obj.get("deck_subtitle") or "")[:200], "slides": slides}

def create_deck(generate, report, case_brief, sources, audit, count=15):
    """One Gemini slide architect call; retry invalid JSON once."""
    catalog = "\n".join(f"[{x['id']}] {x['title']} {x['url']}" for x in sources) or "KHÔNG CÓ NGUỒN WEB ĐƯỢC XÁC THỰC"
    outline = "\n".join(f"{i}. {title}" for i, title in enumerate(ORDER, 1)) if count == 15 else (
        "Business challenge → 4C research → diagnosis → consumer tension → strategic direction "
        "→ Big Idea → IMC execution → budget/KPI → closing."
    )
    prompt = f"""Bạn là SLIDE ARCHITECT cho cuộc thi Marketing Case, không phải essay writer.
Đầu ra là bố cục POWERPOINT 16:9, mỗi slide một insight-led headline, ít chữ, chứng cứ rõ.
Gemini model phân tích đã tạo research report sau 3 vòng phản biện. Hãy sử dụng bản cuối đó.

CASE BRIEF:
{case_brief[:3500]}
FINAL RESEARCH:
{report[:36000]}
EVIDENCE AUDIT:
{audit[:4000]}
DANH MỤC SOURCE IDs HỢP LỆ:
{catalog[:10500]}

Xuất một JSON object CHÍNH XÁC {count} slides (không markdown).
{{"deck_title":"...","deck_subtitle":"...","slides":[
{{"slide_no":1,"section":"...","kicker":"...","headline":"một kết luận cụ thể",
"layout":"cover|split|data|comparison|timeline|statement|framework",
"insights":["ngắn gọn","ngắn gọn","ngắn gọn"],
"visual":"Gợi ý visual cụ thể, không tự bịa số liệu",
"takeaway":"Một câu quyết định chiến lược",
"source_ids":["S1"],"presenter_note":"gợi ý thuyết trình"}}
]}}

BỐ CỤC:
{outline}

NGUYÊN TẮC:
- Đúng {count} slide, đánh số từ 1; headline thể hiện kết luận thay vì title chung chung.
- Mỗi slide 2–4 ý, ngắn gọn và không trùng ý (slide cover có thể 0–2 ý).
- Dùng layout data khi có số liệu thực, comparison cho đối thủ, statement cho tension/Big Idea, timeline cho IMC.
- Mỗi slide phải có Key Takeaway 1 câu (trừ cover có thể trống).
- Giữ đầy đủ mạch: Research → Strategic Diagnosis → Big Idea → IMC → Budget → KPIs.
- Research đã có >=5 insights/mục; trên slide chỉ chọn ý mạnh nhất, không nhồi 5 ý trên một slide.
- Big Idea, tactics, budget allocation và KPI targets là ĐỀ XUẤT chứ không phải fact.
- Khi thiếu dữ liệu, ghi rõ giả thuyết cần kiểm chứng, không tạo bất kỳ số liệu, nguồn hoặc URL mới.
- source_ids chỉ được dùng khi thật sự chứng minh nội dung slide; để [] nếu không chứng minh được.
- Không viết nội dung nào ngoài JSON.
"""
    err = ""
    permitted = {s["id"] for s in sources}
    for attempt in range(2):
        response = generate(prompt + (("\nLỖI LẦN TRƯỚC: " + err + ". Xuất JSON đúng schema.") if attempt else ""))
        try:
            return normalize(response, permitted, count)
        except (ValueError, TypeError, json.JSONDecodeError) as ex:
            err = str(ex)
    raise ValueError("Không tạo được slide deck hợp lệ: " + err)

CSS = """<style>
.case-slide{background:linear-gradient(130deg,#10192e,#1b3357);color:#f0f4fc;
  width:100%;aspect-ratio:16/9;box-sizing:border-box;border-radius:14px;padding:4% 5%;
  display:flex;flex-direction:column;gap:2%;font-family:Arial,sans-serif}
.case-eyebrow{display:flex;justify-content:space-between;font-size:clamp(9px,1.1vw,15px);
  color:#9dc6ee;text-transform:uppercase;letter-spacing:2px;font-weight:bold}
.case-head{font-size:clamp(18px,2.7vw,36px);line-height:1.15;font-weight:800;max-width:100%}
.case-main{display:grid;grid-template-columns:1.2fr .8fr;gap:3%;flex:1;min-height:0}
.case-points{display:flex;flex-direction:column;justify-content:space-evenly;gap:5%}
.case-point{font-size:clamp(10px,1.35vw,20px);display:flex;gap:10px;line-height:1.3}
.case-number{color:#ffda85;font-weight:800}
.case-visual{border-radius:12px;border:1px dashed #8096b1;background:rgba(255,255,255,.05);
  display:flex;flex-direction:column;justify-content:center;align-items:center;
  padding:8%;gap:10%;font-size:clamp(9px,1.14vw,17px);text-align:center}
.case-visual b{color:#ffda85}
.case-foot{border-top:1px solid #657c9d;padding-top:1.5%;font-size:clamp(9px,1.16vw,17px);line-height:1.2}
.case-foot b{color:#ffda85}
.case-refs{font-size:clamp(8px,.83vw,12px);color:#adbfdb}
.case-slide.cover .case-main,.case-slide.statement .case-main{grid-template-columns:1fr}
.case-slide.cover .case-head,.case-slide.statement .case-head{font-size:clamp(23px,3.8vw,55px)}
.case-slide.cover .case-visual,.case-slide.statement .case-visual{display:none}
@media(max-width:550px){.case-visual{display:none}.case-main{grid-template-columns:1fr}}
</style>"""

def slide_html(item):
    esc = html.escape
    bullets = "".join(
        '<div class="case-point"><span class="case-number">' + str(i) + '</span>'
        '<span>' + esc(x) + '</span></div>'
        for i, x in enumerate(item["insights"], 1)
    )
    return (
        '<div class="case-slide '+esc(item["layout"])+'">'
        '<div class="case-eyebrow"><span>'+esc(item["kicker"])+
        '</span><span>'+str(item["slide_no"]).zfill(2)+'</span></div>'
        '<div class="case-head">'+esc(item["headline"])+'</div>'
        '<div class="case-main"><div class="case-points">'+bullets+
        '</div><div class="case-visual"><b>VISUAL / CHART BLUEPRINT</b><span>'+
        esc(item["visual"])+'</span></div></div>'
        '<div class="case-foot"><b>KEY TAKEAWAY — </b>'+esc(item["takeaway"])+'</div>'
        '<div class="case-refs">Citations: '+esc(", ".join(item["source_ids"]) or "Chưa có source ID")+'</div>'
        '</div>'
    )

def ppt_text(slide, msg, x, y, w, h, size=18, bold=False, rgb=(240,244,252)):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(msg)
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*rgb)

def export_pptx(deck, sources):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for item in deck["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(16, 25, 46)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
            Inches(.62), Inches(.5), Inches(.07), Inches(.27))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(255,218,133)
        line.line.fill.background()
        ppt_text(slide, item["kicker"], .85, .48, 10, .4, 11, True, (157,198,238))
        ppt_text(slide, f'{item["slide_no"]:02d}', 12.05, .48, .7, .4, 11, True, (157,198,238))
        ppt_text(slide, item["headline"], .72, 1.12, 11.8, 1.2,
                 29 if len(item["headline"]) < 96 else 22, True)
        wide = item["layout"] in ("cover", "statement")
        for j, point in enumerate(item["insights"][:4]):
            ppt_text(slide, f'{j+1}. {point}', .85, 2.48+j*.7, 11.2 if wide else 6.4,
                     .65, 19 if wide else 17)
        if not wide:
            panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(7.8), Inches(2.42), Inches(4.7), Inches(3.08))
            panel.fill.solid()
            panel.fill.fore_color.rgb = RGBColor(28, 49, 83)
            panel.line.color.rgb = RGBColor(97, 125, 163)
            ppt_text(slide, "VISUAL / CHART BLUEPRINT", 8.14, 2.71, 4, .45,
                     12, True, (255,218,133))
            ppt_text(slide, item["visual"], 8.14, 3.37, 4.0, 1.65, 16)
        ppt_text(slide, "KEY TAKEAWAY", .75, 6.00, 2.0, .35, 11, True, (255,218,133))
        ppt_text(slide, item["takeaway"], 2.65, 5.94, 9.75, .77, 16, True)
        ppt_text(slide, "Sources: " + ", ".join(item["source_ids"]), .75, 7.1, 11.5, .26,
                 9, False, (157,198,238))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

def _navigate_slide(delta, total_slides):
    """Streamlit callbacks run before widgets are instantiated on the next rerun."""
    current = st.session_state.get("slide_choice", 0)
    st.session_state["slide_choice"] = max(0, min(total_slides - 1, current + delta))


def show(deck, sources):
    slides = deck["slides"]
    if not slides:
        st.warning("Deck chưa có slide nào.")
        return
    # When a new case has a different slide count, reset out-of-range selection BEFORE selectbox.
    if (not isinstance(st.session_state.get("slide_choice"), int)
            or st.session_state.get("slide_choice", 0) not in range(len(slides))):
        st.session_state["slide_choice"] = 0
    st.caption(f"{len(slides)} editable slides · 16:9 · Strategic case pitch")
    n = st.selectbox("Slide", range(len(slides)), key="slide_choice",
         format_func=lambda i: f"{i+1:02d} — {slides[i]['section']}")
    st.markdown(CSS, unsafe_allow_html=True)
    st.markdown(slide_html(slides[n]), unsafe_allow_html=True)
    col_a, col_b, col_c = st.columns([1, 2, 1])
    with col_a:
        st.button("← Trước", disabled=(n==0),
                  on_click=_navigate_slide, args=(-1, len(slides)))
    with col_b:
        st.caption(f"Slide {n+1} / {len(slides)} · Layout {slides[n]['layout']}")
    with col_c:
        st.button("Tiếp →", disabled=(n==len(slides)-1),
                  on_click=_navigate_slide, args=(1, len(slides)))
    with st.expander("Speaker notes & sources"):
        st.write(slides[n]["presenter_note"])
        for sid in slides[n]["source_ids"]:
            source = next((s for s in sources if s["id"] == sid), None)
            if source:
                st.markdown(f"[{sid}] [{source['title']}]({source['url']})")
    st.caption("Visual blueprint là gợi ý thiết kế, không phải biểu đồ dữ liệu đã xác minh.")
    st.download_button("⬇️ PowerPoint 16:9 (chỉnh sửa được)",
         data=export_pptx(deck, sources), file_name="marketing_case_solver.pptx",
         mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
         use_container_width=True)
    st.download_button("⬇️ Slide JSON", json.dumps(deck,ensure_ascii=False,indent=2),
         "marketing_case_slides.json", "application/json", use_container_width=True)
